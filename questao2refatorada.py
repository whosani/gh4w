from openpyxl import load_workbook
arquivo_excel = load_workbook('base_pedidos_v2 1.xlsx')
pedidos = arquivo_excel['Planilha1']
dicionario = {}
numero_linhas = pedidos.max_row + 1
for linha in range(2, numero_linhas):
    coluna_data_entrega = pedidos.cell(row=linha, column=5).value
    if coluna_data_entrega != None:
        id_cliente = pedidos.cell(row=linha, column=2).value
        qntd_produto = pedidos.cell(row=linha, column=7).value
        valor = dicionario.get(id_cliente)
        if valor != None:
            dicionario[id_cliente] = valor + qntd_produto
        else: 
            dicionario[id_cliente] = qntd_produto
print(dicionario)

